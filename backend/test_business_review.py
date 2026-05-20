"""一键测试：创建业务复盘会测试数据 → 生成 PPTX"""
import sys, os, asyncio, tempfile
sys.path.insert(0, os.path.dirname(__file__))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl import Workbook

OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "test_output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "机构对照数据"

headers = [
    "机构名称", "当月保费(万元)", "保费达成率", "承保件数", "活动率",
    "队伍人数", "人均产能(万元)", "13个月继续率", "25个月继续率",
    "投产比", "同比增长率", "荣誉达成", "问题描述", "改进措施"
]
ws.append(headers)

header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="1a365d", end_color="1a365d", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
for cell in ws[1]:
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

agencies = [
    ["深圳福田支公司", 2350.80, "118.6%", 480, "82.3%", 55, 42.74, "96.5%", "93.2%", 0.72, "18.9%",
     "IDA x3, MDRT x2, MCI x1", "", ""],
    ["北京朝阳支公司", 1850.50, "105.2%", 360, "78.5%", 48, 38.55, "95.2%", "91.8%", 0.85, "12.3%",
     "IDA x2, MDRT x1", "高净值客户开拓不足", "加强私行渠道合作, 每月2场高端沙龙"],
    ["成都锦江支公司", 1450.60, "98.5%", 290, "72.8%", 42, 34.54, "94.1%", "90.3%", 0.95, "6.8%",
     "MDRT x1", "25J继续率连续2月下滑", "加强客户维系, 续期提醒前置30天"],
    ["杭州西湖支公司", 1280.30, "92.0%", 255, "70.1%", 38, 33.69, "93.8%", "89.5%", 1.05, "3.2%",
     "IDA x1", "队伍新增人力培训周期过长", "优化新人90天成长营, 缩短独立展业周期"],
    ["广州天河支公司", 1100.20, "85.1%", 220, "65.2%", 35, 31.43, "92.1%", "88.5%", 1.18, "-3.5%",
     "", "活动率持续走低, 主管缺位严重", "每日活动量追踪, 主管每周1次陪访"],
    ["上海浦东支公司", 980.50, "82.3%", 195, "68.5%", 36, 27.24, "90.5%", "86.2%", 1.25, "-5.8%",
     "", "客均产能偏低, 件均保费下滑", "推动产品组合销售, 培训需求导向面谈技巧"],
    ["武汉洪山支公司", 890.30, "78.2%", 180, "62.1%", 32, 27.82, "91.2%", "87.8%", 1.32, "-8.1%",
     "", "团队流失严重, 2名主管离职", "启动主管储备计划, 完善离职人员客户交接流程"],
    ["重庆渝北支公司", 820.60, "75.5%", 165, "60.8%", 30, 27.35, "89.8%", "85.6%", 1.40, "-9.2%",
     "", "新增保费缺口达180万, 补量困难", "突击老客户加保, 推动团险交叉销售"],
    ["南京鼓楼支公司", 720.40, "68.2%", 150, "55.4%", 28, 25.73, "88.5%", "84.1%", 1.52, "-11.5%",
     "", "市场竞品冲击严重, 主力产品竞争力下降", "加大产品培训力度, 推动组合拳销售策略"],
    ["郑州金水支公司", 580.80, "58.5%", 128, "52.1%", 25, 23.23, "87.3%", "83.2%", 1.68, "-15.2%",
     "", "新人留存率仅35%, 队伍断层严重", "暂停盲目增员, 聚焦在职人员技能训练"],
]

data_font = Font(name="微软雅黑", size=10)
data_align = Alignment(vertical="center", wrap_text=True)
green_fill = PatternFill(start_color="dcfce7", end_color="dcfce7", fill_type="solid")
yellow_fill = PatternFill(start_color="fef9c3", end_color="fef9c3", fill_type="solid")
red_fill = PatternFill(start_color="fee2e2", end_color="fee2e2", fill_type="solid")

for ri, row_data in enumerate(agencies):
    ws.append(row_data)
    row_num = ri + 2
    for ci in range(1, len(row_data) + 1):
        cell = ws.cell(row=row_num, column=ci)
        cell.font = data_font
        cell.alignment = data_align
        cell.border = thin_border

    rate_str = str(row_data[2])
    try:
        rate_val = float(rate_str.rstrip("%"))
    except ValueError:
        rate_val = 0
    if rate_val >= 100:
        fill = green_fill
    elif rate_val >= 80:
        fill = yellow_fill
    else:
        fill = red_fill
    for ci in range(1, 4):
        ws.cell(row=row_num, column=ci).fill = fill

column_widths = [18, 15, 12, 10, 10, 10, 15, 14, 14, 10, 12, 25, 35, 35]
for i, w in enumerate(column_widths):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w

for ri in range(2, 12):
    ws.row_dimensions[ri].height = 28

excel_path = os.path.join(OUTPUT_DIR, "业务对标复盘_测试数据.xlsx")
wb.save(excel_path)
print(f"[1/3] 测试 Excel 已生成: {excel_path}")
print(f"      10个机构, {len(headers)}项指标, 含🟢🟡🔴状态着色")

async def generate_pptx():
    from slide_builder.pipeline import GenerationPipeline

    pipeline = GenerationPipeline(
        scene="insurance",
        template_id="professional-blue",
        ai_mode="offline",
        meeting_type="business_review",
        excel_filepath=excel_path,
    )
    result = await pipeline.run("")

    pptx_path = os.path.join(OUTPUT_DIR, "业务对标复盘_生成结果.pptx")
    with open(pptx_path, "wb") as f:
        f.write(result.pptx_bytes)

    print(f"\n[2/3] PPTX 已生成: {pptx_path}")
    print(f"      模式: {result.mode}")
    print(f"      页面数: {len(result.slides)}")
    print(f"      标题: {result.title}")
    print(f"      信息: {result.message}")
    if result.qa_results:
        errors = [q for q in result.qa_results if q.get("severity") == "error"]
        warns = [q for q in result.qa_results if q.get("severity") == "warning"]
        if errors:
            print(f"      ❌ 错误: {len(errors)}条")
        if warns:
            print(f"      ⚠️ 警告: {len(warns)}条")
        print(f"      ℹ️ 校验: {len(result.qa_results)}条")

    print(f"\n[3/3] ✅ 完成! 请打开以下文件查看效果:")
    print(f"      {pptx_path}")
    return result

result = asyncio.run(generate_pptx())