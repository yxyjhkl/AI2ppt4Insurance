import openpyxl
from openpyxl.utils import get_column_letter
from typing import List, Dict, Any, Optional, Tuple
from dataclasses import dataclass, field
import re
import statistics


@dataclass
class ColumnMeta:
    name: str
    index: int
    letter: str
    data_type: str
    sample_values: List[Any] = field(default_factory=list)
    null_count: int = 0
    stats: Dict[str, Any] = field(default_factory=dict)


@dataclass
class SheetData:
    name: str
    headers: List[str]
    columns: List[ColumnMeta]
    rows: List[List[Any]]
    row_count: int
    col_count: int


@dataclass
class ExcelData:
    filename: str
    sheets: List[SheetData]
    total_rows: int = 0
    errors: List[str] = field(default_factory=list)


def parse_excel(filepath: str, header_row: int = 0) -> ExcelData:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    filename = filepath.rsplit("\\", 1)[-1].rsplit("/", 1)[-1]
    sheets = []
    total_rows = 0
    errors = []

    for ws in wb.worksheets:
        try:
            sheet = _parse_sheet(ws, header_row)
            sheets.append(sheet)
            total_rows += sheet.row_count
        except Exception as e:
            errors.append(f"Sheet '{ws.title}': {e}")

    wb.close()
    return ExcelData(filename=filename, sheets=sheets, total_rows=total_rows, errors=errors)


def _parse_sheet(ws, header_row: int) -> SheetData:
    merged_map = {}
    for merged_range in ws.merged_cells.ranges:
        min_col = merged_range.min_col
        min_row = merged_range.min_row
        max_col = merged_range.max_col
        max_row = merged_range.max_row
        top_left_value = ws.cell(row=min_row, column=min_col).value
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                merged_map[(row, col)] = top_left_value

    all_rows = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        row_data = []
        for cell in row:
            val = cell.value
            if val is None and (cell.row, cell.column) in merged_map:
                val = merged_map[(cell.row, cell.column)]
            row_data.append(val)
        all_rows.append(row_data)

    if not all_rows:
        return SheetData(name=ws.title, headers=[], columns=[], rows=[], row_count=0, col_count=0)

    header_idx = max(0, min(header_row, len(all_rows) - 1))
    raw_headers = [str(h) if h is not None else f"Column_{i+1}" for i, h in enumerate(all_rows[header_idx])]

    headers = []
    seen = {}
    for h in raw_headers:
        h_clean = h.strip()
        if h_clean in seen:
            seen[h_clean] += 1
            headers.append(f"{h_clean}_{seen[h_clean]}")
        else:
            seen[h_clean] = 0
            headers.append(h_clean)

    data_rows = all_rows[header_idx + 1:]
    clean_rows = []
    for row in data_rows:
        clean = [v for v in row] + [None] * (len(headers) - len(row))
        clean_rows.append(clean[:len(headers)])

    columns = []
    for ci, header in enumerate(headers):
        col_letter = get_column_letter(ci + 1)
        values = [row[ci] for row in clean_rows if ci < len(row)]
        non_null = [v for v in values if v is not None and str(v).strip()]
        null_count = len(values) - len(non_null)
        data_type, sample = _infer_type(non_null)
        stats = _compute_stats(non_null, data_type)
        columns.append(ColumnMeta(
            name=header, index=ci, letter=col_letter,
            data_type=data_type, sample_values=sample,
            null_count=null_count, stats=stats
        ))

    return SheetData(
        name=ws.title, headers=headers, columns=columns,
        rows=clean_rows, row_count=len(clean_rows), col_count=len(headers)
    )


def _infer_type(values: List[Any]) -> Tuple[str, List[Any]]:
    if not values:
        return "empty", []
    numeric = []
    percent = []
    text_val = []
    for v in values:
        s = str(v).strip()
        if s.endswith("%"):
            try:
                float(s.rstrip("%"))
                percent.append(v)
                continue
            except ValueError:
                pass
        try:
            float(s.replace(",", "").replace("，", ""))
            numeric.append(v)
        except (ValueError, TypeError):
            text_val.append(v)

    if len(percent) > len(numeric) and len(percent) > len(text_val):
        return "percent", percent[:5]
    if len(numeric) > len(text_val):
        return "number", numeric[:5]
    date_count = sum(1 for v in text_val if re.search(r'\d{4}[-/年]\d{1,2}[-/月]\d{1,2}', str(v)))
    if date_count > len(text_val) * 0.5:
        return "date", text_val[:5]
    return "text", text_val[:5]


def _compute_stats(values: List[Any], data_type: str) -> Dict[str, Any]:
    if data_type not in ("number", "percent"):
        return {}
    nums = []
    for v in values:
        try:
            s = str(v).strip().rstrip("%")
            n = float(s.replace(",", "").replace("，", ""))
            nums.append(n)
        except (ValueError, TypeError):
            pass
    if not nums:
        return {}
    return {
        "min": round(min(nums), 2),
        "max": round(max(nums), 2),
        "avg": round(statistics.mean(nums), 2),
        "median": round(statistics.median(nums), 2),
        "count": len(nums),
        "sum": round(sum(nums), 2),
    }


def to_prompt_summary(data: ExcelData, max_rows_per_sheet: int = 20) -> str:
    lines = [f"Excel文件: {data.filename}", f"共 {len(data.sheets)} 个工作表, {data.total_rows} 行数据", ""]
    for sheet in data.sheets:
        lines.append(f"--- 工作表: {sheet.name} ({sheet.row_count}行 × {sheet.col_count}列) ---")
        lines.append(f"列名: {' | '.join(sheet.headers)}")
        for col in sheet.columns:
            type_info = col.data_type
            if col.stats:
                stats_str = ", ".join(f"{k}={v}" for k, v in col.stats.items())
                type_info += f" [{stats_str}]"
            lines.append(f"  {col.letter}: {col.name} ({type_info})")
        lines.append(f"数据预览 (前{min(max_rows_per_sheet, sheet.row_count)}行):")
        lines.append(" | ".join(sheet.headers))
        lines.append(" | ".join("---" for _ in sheet.headers))
        for row in sheet.rows[:max_rows_per_sheet]:
            lines.append(" | ".join(str(v) if v is not None else "" for v in row))
        lines.append("")
    return "\n".join(lines)