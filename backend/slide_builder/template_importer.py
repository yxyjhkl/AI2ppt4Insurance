"""Template importer - extracts theme/colors/fonts from existing PPTX files."""
from pptx import Presentation
import json
import os


def import_template_from_pptx(pptx_path: str, output_dir: str) -> dict:
    prs = Presentation(pptx_path)

    theme = {
        "colors": _extract_colors(prs),
        "fonts": _extract_fonts(prs),
        "slide_width": prs.slide_width,
        "slide_height": prs.slide_height,
        "slide_count": len(prs.slides),
    }

    theme_path = os.path.join(output_dir, "theme.json")
    with open(theme_path, "w", encoding="utf-8") as f:
        json.dump(theme, f, ensure_ascii=False, indent=2)

    return theme


def _extract_colors(prs: Presentation) -> dict:
    colors = {
        "primary": "1e40af",
        "secondary": "3b82f6",
        "accent": "f59e0b",
        "background": "FFFFFF",
        "text": "1f2937",
    }

    if prs.slide_master and prs.slide_master.slide_layouts:
        for layout in prs.slide_master.slide_layouts:
            if layout.placeholders:
                for ph in layout.placeholders:
                    if ph.has_text_frame:
                        for para in ph.text_frame.paragraphs:
                            for run in para.runs:
                                if run.font.color and run.font.color.rgb:
                                    hex_val = str(run.font.color.rgb)
                                    key = "primary" if para.font.size and para.font.size > 200000 else "text"
                                    colors[key] = hex_val

    return colors


def _extract_fonts(prs: Presentation) -> dict:
    fonts = {"title": "Arial", "body": "Calibri"}

    if prs.slide_master:
        try:
            font = prs.slide_master.element.find(
                "{http://schemas.openxmlformats.org/drawingml/2006/main}titleFont"
            )
            if font is not None and font.get("typeface"):
                fonts["title"] = font.get("typeface")
        except Exception:
            pass

    return fonts
