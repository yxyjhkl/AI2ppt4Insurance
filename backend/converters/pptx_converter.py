"""PPTX to Markdown converter - extracts text, structure, and notes from presentations."""
import re
from typing import Optional


def pptx_to_markdown(filepath: str) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        raise RuntimeError("python-pptx 未安装，无法解析 PPTX 文件")

    prs = Presentation(filepath)
    lines: list[str] = []

    for i, slide in enumerate(prs.slides, 1):
        title_text = ""
        body_texts: list[str] = []
        table_texts: list[str] = []
        notes_text = ""

        for shape in slide.shapes:
            if shape.is_placeholder:
                ph = shape.placeholder_format
                if ph.type == 1:
                    title_text = _extract_shape_text(shape)
                elif ph.type in (2, 7):
                    body_texts.append(_extract_shape_text(shape))
            elif shape.has_text_frame:
                body_texts.append(_extract_shape_text(shape))

            if shape.has_table:
                table_texts.append(_extract_table_md(shape))

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        lines.append(f"## 第{i}页: {title_text}" if title_text else f"## 第{i}页")
        for bt in body_texts:
            text = bt.strip()
            if text:
                for line in text.split("\n"):
                    stripped = line.strip()
                    if stripped:
                        lines.append(f"- {stripped}")
        for tt in table_texts:
            lines.append("")
            lines.append(tt)
        if notes_text:
            lines.append(f"> 备注: {notes_text}")
        lines.append("")

    return "\n".join(lines)


def _extract_shape_text(shape) -> str:
    if not shape.has_text_frame:
        return ""
    return shape.text_frame.text


def _extract_table_md(shape) -> str:
    table = shape.table
    rows: list[list[str]] = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append(cells)
    if not rows:
        return ""
    max_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_cols:
            r.append("")

    md_lines: list[str] = []
    header = "| " + " | ".join(rows[0]) + " |"
    sep = "|" + "|".join("---" for _ in rows[0]) + "|"
    md_lines.append(header)
    md_lines.append(sep)
    for row in rows[1:]:
        md_lines.append("| " + " | ".join(row) + " |")
    return "\n".join(md_lines)
